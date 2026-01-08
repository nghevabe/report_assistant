from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.cell.cell import MergedCell
from copy import copy


def find_last_column_by_header(ws):
    """
    Trả về index của cột cuối cùng có ít nhất 1 cell có value
    """
    max_col = ws.max_column
    max_row = ws.max_row

    for col in range(max_col, 0, -1):          # duyệt ngược
        for row in range(1, max_row + 1):
            cell = ws.cell(row=row, column=col)
            if cell.value not in (None, ""):
                return col

    return 1


def copy_excel_columns(
    input_file: str,
    output_file: str,
    sheet_name: str | None = None,
    src_start_col_letter: str = "AC",
    header_row: int = 1,
    gap_cols: int = 5,   # khoảng cách giữa block cũ và block paste,
    matrix_revenue=None
):
    if matrix_revenue is None:
        matrix_revenue = []
    wb = load_workbook(input_file)
    ws = wb[sheet_name] if sheet_name else wb.active

    src_start_col = column_index_from_string(src_start_col_letter)
    src_end_col = find_last_column_by_header(ws)

    # vùng cần copy
    src_width = src_end_col - src_start_col + 1

    # paste bắt đầu sau vùng source + gap
    dest_start_col = src_end_col + gap_cols
    dest_end_col = dest_start_col + src_width - 1

    max_row = ws.max_row

    print(f"Copy columns {src_start_col_letter} → {get_column_letter(src_end_col)}")
    print(f"Paste from column {get_column_letter(dest_start_col)} → {get_column_letter(dest_end_col)}")

    # ===============================
    # 1) Copy column width
    # ===============================
    for col in range(src_start_col, src_end_col + 1):
        src_letter = get_column_letter(col)
        dest_letter = get_column_letter(dest_start_col + (col - src_start_col))
        ws.column_dimensions[dest_letter].width = ws.column_dimensions[src_letter].width

    # ===============================
    # 2) Copy cell value + style (skip merged read-only cells)
    # ===============================
    for row in range(1, max_row + 1):
        for col in range(src_start_col, src_end_col + 1):
            src_cell = ws.cell(row=row, column=col)
            dest_cell = ws.cell(row=row, column=dest_start_col + (col - src_start_col))

            # Skip non-top-left merged cells (read-only)
            if isinstance(src_cell, MergedCell) or isinstance(dest_cell, MergedCell):
                continue

            dest_cell.value = src_cell.value

            if src_cell.has_style:
                dest_cell.font = copy(src_cell.font)
                dest_cell.border = copy(src_cell.border)
                dest_cell.fill = copy(src_cell.fill)
                dest_cell.number_format = copy(src_cell.number_format)
                dest_cell.alignment = copy(src_cell.alignment)
                dest_cell.protection = copy(src_cell.protection)

    # ===============================
    # 3) Copy merged ranges (FIX: iterate over copy)
    # ===============================
    merged_ranges = list(ws.merged_cells.ranges)

    for merged_range in merged_ranges:
        min_col, min_row, max_col, max_row2 = merged_range.bounds

        # chỉ copy các merge nằm trong vùng source
        if min_col >= src_start_col and max_col <= src_end_col:
            new_min_col = dest_start_col + (min_col - src_start_col)
            new_max_col = dest_start_col + (max_col - src_start_col)

            ws.merge_cells(
                start_row=min_row,
                start_column=new_min_col,
                end_row=max_row2,
                end_column=new_max_col,
            )

    row = 3

    first_col_index = dest_start_col

    write_table_with_matrix(output_file, ws, wb, matrix_revenue, dest_start_col)


def write_table_with_matrix(output_file, ws, wb, matrix_revenue, dest_start_col):

    col_index = 0
    for col_item in matrix_revenue:
        row_index = 0
        for row_item in col_item:
            revenue_value = row_item.split("=")[1]
            ws.cell(row=3 + row_index, column=dest_start_col + col_index).value = int(revenue_value)
            row_index += 1

        col_index += 1

    wb.save(output_file)
    print(f"✅ Done. Output saved to: {output_file}")
